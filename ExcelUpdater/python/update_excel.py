#!/usr/bin/env python3
"""
Update cells in an Excel workbook on a network share using openpyxl.

Formulas are NOT calculated here. The workbook is flagged so Excel recalculates
everything the next time the file is opened.

Exit codes (for AutoSys): 0 = success, 1 = config/validation, 2 = file/IO, 3 = unexpected.
"""
import argparse
import io
import json
import sys
import time
import traceback
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

EXIT_OK, EXIT_CONFIG, EXIT_IO, EXIT_FAILURE = 0, 1, 2, 3


class ConfigError(Exception):
    pass


def log(level: str, msg: str) -> None:
    print(f"{datetime.now():%Y-%m-%d %H:%M:%S} [{level}] {msg}", flush=True)


def load_config(path: Path) -> dict:
    try:
        cfg = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as ex:
        raise ConfigError(f"Cannot read config {path}: {ex}") from ex

    if not cfg.get("filePath"):
        raise ConfigError("filePath is required.")
    if not cfg.get("updates"):
        raise ConfigError("At least one entry in updates is required.")
    for u in cfg["updates"]:
        if not u.get("sheet") or not u.get("cell"):
            raise ConfigError("Every update needs sheet and cell.")

    cfg.setdefault("backupFolder", None)
    cfg.setdefault("backupRetentionDays", 30)
    cfg.setdefault("lockRetryCount", 6)
    cfg.setdefault("lockRetryDelaySeconds", 10)
    return cfg


def open_exclusive(cfg: dict):
    """Open for read/write. Fails if another user (e.g. Excel) has the file open."""
    path = Path(cfg["filePath"])
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    attempts = cfg["lockRetryCount"]
    for attempt in range(1, attempts + 1):
        try:
            return open(path, "r+b")
        except OSError as ex:
            if attempt == attempts:
                raise
            log("WARN", f"File is locked (attempt {attempt}/{attempts}): {ex}. "
                        f"Retrying in {cfg['lockRetryDelaySeconds']}s.")
            time.sleep(cfg["lockRetryDelaySeconds"])


def backup(cfg: dict, data: bytes) -> None:
    folder = cfg["backupFolder"]
    if not folder:
        return

    src = Path(cfg["filePath"])
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    dest = folder / f"{src.stem}_{datetime.now():%Y%m%d_%H%M%S}{src.suffix}"
    dest.write_bytes(data)
    log("INFO", f"Backup written: {dest}")

    days = cfg["backupRetentionDays"]
    if days <= 0:
        return
    cutoff = time.time() - days * 86400
    for old in folder.glob(f"{src.stem}_*{src.suffix}"):
        try:
            if old.stat().st_mtime < cutoff:
                old.unlink()
                log("INFO", f"Deleted old backup: {old}")
        except OSError as ex:
            log("WARN", f"Could not delete {old}: {ex}")


def convert(update: dict):
    raw = str(update.get("value", ""))
    kind = update.get("type", "string").lower()

    if kind == "number":
        number = float(raw)
        return int(number) if number.is_integer() and "." not in raw else number
    if kind == "date":
        if raw == "{today}":
            return datetime.combine(datetime.today().date(), datetime.min.time())
        return datetime.strptime(raw, "%Y-%m-%d")
    if kind == "bool":
        if raw.lower() not in ("true", "false"):
            raise ConfigError(f"Bad bool '{raw}' for {update['sheet']}!{update['cell']}.")
        return raw.lower() == "true"
    if kind == "string":
        return raw
    raise ConfigError(f"Unknown type '{kind}' for {update['sheet']}!{update['cell']} "
                      f"(use number, date, bool, string).")


def apply_updates(wb, cfg: dict) -> dict:
    expected = {}
    for u in cfg["updates"]:
        if u["sheet"] not in wb.sheetnames:
            raise ConfigError(f"Worksheet '{u['sheet']}' not found.")
        cell = wb[u["sheet"]][u["cell"]]
        if cell.data_type == "f":
            raise ConfigError(f"{u['sheet']}!{u['cell']} contains a formula; refusing to overwrite it.")

        new_value = convert(u)
        log("INFO", f"{u['sheet']}!{u['cell']}: '{cell.value}' -> '{new_value}'")
        cell.value = new_value
        expected[(u["sheet"], u["cell"])] = new_value
    return expected


def count_formulas(wb) -> int:
    return sum(1 for ws in wb.worksheets for row in ws.iter_rows()
               for c in row if c.data_type == "f")


def run(cfg: dict) -> None:
    log("INFO", f"Starting. Target: {cfg['filePath']}")

    with open_exclusive(cfg) as fh:
        original = fh.read()
        backup(cfg, original)

        # Formulas are loaded as text (data_only=False) and written back unchanged.
        wb = load_workbook(io.BytesIO(original))
        formulas_before = count_formulas(wb)

        expected = apply_updates(wb, cfg)

        # openpyxl cannot calculate. Ask Excel to do a full recalculation on next open.
        wb.calculation.fullCalcOnLoad = True

        out = io.BytesIO()
        wb.save(out)
        payload = out.getvalue()

        # Sanity check the result before touching the real file.
        check = load_workbook(io.BytesIO(payload))
        formulas_after = count_formulas(check)
        if formulas_after != formulas_before:
            raise RuntimeError(f"Formula count changed ({formulas_before} -> {formulas_after}); file left unchanged.")
        for (sheet, addr), value in expected.items():
            if check[sheet][addr].value != value:
                raise RuntimeError(f"Verification failed for {sheet}!{addr}; file left unchanged.")

        # Only now write back to the open handle.
        fh.seek(0)
        fh.truncate()
        fh.write(payload)
        fh.flush()

    log("INFO", f"File updated successfully ({formulas_after} formulas preserved, recalculated on next open).")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", default=str(Path(__file__).with_name("config.json")))
    parser.add_argument("--file", help="Override filePath from config")
    args = parser.parse_args()

    try:
        cfg = load_config(Path(args.config))
        if args.file:
            cfg["filePath"] = args.file
        run(cfg)
        return EXIT_OK
    except ConfigError as ex:
        log("ERROR", str(ex))
        return EXIT_CONFIG
    except OSError as ex:  # includes PermissionError, FileNotFoundError
        log("ERROR", f"File access problem: {ex}")
        return EXIT_IO
    except Exception:
        log("ERROR", f"Unexpected failure:\n{traceback.format_exc()}")
        return EXIT_FAILURE


if __name__ == "__main__":
    sys.exit(main())
